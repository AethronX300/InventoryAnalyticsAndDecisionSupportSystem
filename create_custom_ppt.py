import collections 
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# Theme Colors
BG_COLOR = RGBColor(19, 28, 23)        # #131c17
CARD_COLOR = RGBColor(26, 36, 30)      # #1a241e
TEXT_MAIN = RGBColor(233, 242, 234)    # #e9f2ea
TEXT_SUB = RGBColor(125, 148, 133)     # #7d9485
ACCENT = RGBColor(127, 180, 141)       # #7fb48d

prs = Presentation()
# Set slide dimensions to 16:9
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)

blank_slide_layout = prs.slide_layouts[6] # Blank layout

def set_background(slide):
    # Add a full screen rectangle for background
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_COLOR
    bg.line.color.rgb = BG_COLOR
    
    # Add a decorative accent line at the top
    top_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.1))
    top_line.fill.solid()
    top_line.fill.fore_color.rgb = ACCENT
    top_line.line.color.rgb = ACCENT

def create_title_slide(title_text, subtitle_text):
    slide = prs.slides.add_slide(blank_slide_layout)
    set_background(slide)
    
    # Title Box
    tx_box = slide.shapes.add_textbox(Inches(1), Inches(2.5), prs.slide_width - Inches(2), Inches(1.5))
    tf = tx_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.bold = True
    p.font.size = Pt(54)
    p.font.color.rgb = TEXT_MAIN
    p.alignment = PP_ALIGN.CENTER
    
    # Subtitle Box
    sub_box = slide.shapes.add_textbox(Inches(1), Inches(4.5), prs.slide_width - Inches(2), Inches(2))
    tf_sub = sub_box.text_frame
    tf_sub.word_wrap = True
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = subtitle_text
    p_sub.font.size = Pt(28)
    p_sub.font.color.rgb = TEXT_SUB
    p_sub.alignment = PP_ALIGN.CENTER

def create_content_slide(title_text, bullet_points):
    slide = prs.slides.add_slide(blank_slide_layout)
    set_background(slide)
    
    # Title Box
    tx_box = slide.shapes.add_textbox(Inches(1), Inches(0.5), prs.slide_width - Inches(2), Inches(1))
    tf = tx_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.bold = True
    p.font.size = Pt(40)
    p.font.color.rgb = TEXT_MAIN
    
    # Divider line under title
    div = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1.5), Inches(2.5), Inches(0.05))
    div.fill.solid()
    div.fill.fore_color.rgb = ACCENT
    div.line.color.rgb = ACCENT
    
    # Content Card Background
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(2), prs.slide_width - Inches(2), prs.slide_height - Inches(2.5))
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_COLOR
    card.line.color.rgb = CARD_COLOR
    
    # Content Text Box
    content_box = slide.shapes.add_textbox(Inches(1.5), Inches(2.5), prs.slide_width - Inches(3), prs.slide_height - Inches(3.5))
    tf_content = content_box.text_frame
    tf_content.word_wrap = True
    
    for i, point in enumerate(bullet_points):
        p = tf_content.add_paragraph() if i > 0 else tf_content.paragraphs[0]
        p.text = f"•  {point}"
        p.font.size = Pt(28)
        p.font.color.rgb = TEXT_MAIN
        p.space_after = Pt(20)

# 1. Title
create_title_slide(
    "Inventory Analytics & Decision Support System",
    "Name: [Your Name]\nRoll No: [Your Roll No]\nSubject: [Subject]\nInstitute: [Institute Name]"
)

# 2. Intro
create_content_slide("Introduction", [
    "A dynamic, real-time inventory management and analytics dashboard.",
    "Purpose: Replaces manual tracking with an automated, data-driven approach.",
    "Built using: Python, Dash (React), and SQLite for a highly responsive interface."
])

# 3. Features
create_content_slide("Key Features & Capabilities", [
    "Real-Time Analytics: Instant visibility into revenue and units sold.",
    "Inventory Health: Tracks products by stock value and identifies low-stock items.",
    "Checkout Terminal: Dedicated Point of Sale (POS) for recording transactions.",
    "Decision Support: Flags dead stock and automatically identifies demand surges."
])

# 4. Architecture
create_content_slide("Technical Architecture", [
    "Frontend UI: Developed using Dash and Plotly to create an interactive SPA.",
    "Backend Logic: Python-based data processing with fast in-memory caching.",
    "Database: Deterministically seeded SQLite database for robust local storage.",
    "Visualizations: Dynamic graphs with unified theming and responsive callbacks."
])

# 5. Business Value
create_content_slide("Business Value & Impact", [
    "Loss Prevention: Immediate alerts for critical stockouts prevent lost revenue.",
    "Optimized Purchasing: Built-in formulas for tracking cover days and restocking.",
    "Customizable: Adjust currency settings and policy thresholds (e.g. surge %).",
    "Data-Driven Decisions: Intuitive KPI cards instantly guide business strategy."
])

# 6. Conclusion
create_content_slide("Conclusion & Future Scope", [
    "Conclusion: The system successfully bridges the gap between raw sales data and actionable business intelligence.",
    "Future Scope 1: Integration with external payment gateways.",
    "Future Scope 2: Connecting with supplier APIs for automated inventory replenishment.",
    "Future Scope 3: Implementing predictive AI machine learning models for demand forecasting."
])

prs.save('Meridian_Themed_Presentation.pptx')
print('Themed PPT generated as Meridian_Themed_Presentation.pptx')
